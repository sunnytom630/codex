"""官方数据提取。保留年度代码、原文项目名称和来源；OCR数字须通过累计校验。"""
from pathlib import Path
import sys,json,re,hashlib,collections,csv
sys.path.insert(0,'/private/tmp/hebei-data-libs')
import openpyxl,xlrd
ROOT=Path(__file__).parent
sources=[]
def add_source(id,title,url,path,published,kind,publisher='河北省教育考试院'):
 p=ROOT/'data/raw'/path
 sources.append({'id':id,'title':title,'url':url,'publisher':publisher,'published':published,'checked':'2026-10-29','kind':kind,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
cutoffs=[]
urls={
 '2023-history':'https://file.hebeea.edu.cn/files/article/2023/07/20230725082954_71.xlsx',
 '2023-physics':'https://file.hebeea.edu.cn/files/article/2023/07/20230725082954_586.xlsx',
 '2026-history':'https://file.hebeea.edu.cn/upload/resources/file/2026/07/22/27221.xlsx',
 '2026-physics':'https://file.hebeea.edu.cn/upload/resources/file/2026/07/22/27220.xlsx',
 '2025-history':'https://file.hebeea.edu.cn/files/article/2025/07/20250722214851_332.xlsx',
 '2025-physics':'https://file.hebeea.edu.cn/files/article/2025/07/20250722214852_210.xlsx',
 '2024-history':'https://news.huanbohainews.com.cn/resource/2024-07/22/08723adb-64c3-41b7-83c5-9ad81fad4300.xlsx',
 '2024-physics':'https://news.huanbohainews.com.cn/resource/2024-07/22/5b01f05d-a30a-4c4b-bd99-12db7833c520.xlsx'}
for stem,url in urls.items():
 year=int(stem[:4]);category='历史' if stem.endswith('history') else '物理'
 path=ROOT/'data/raw'/f'{stem}.xlsx';w=openpyxl.load_workbook(path,read_only=True,data_only=True);s=w.active
 assert f'{year}年河北省普通高校招生本科批' in s.cell(1,1).value
 assert category in s.cell(2,7).value
 add_source(stem,f'{year}河北本科批{category}组合平行志愿投档统计',url,path.name,f'{year}-07-25' if year==2023 else f'{year}-07-22','cutoff')
 for rownum,r in enumerate(s.iter_rows(min_row=6,values_only=True),6):
  if not isinstance(r[0],str) or not re.fullmatch(r'\d{4}',r[0]) or not r[1] or not r[3]:continue
  score=r[4] if r[4] not in (None,'') else None
  if score is not None and (not isinstance(score,(int,float)) or score!=int(score) or not 100<=score<=750):raise ValueError((stem,rownum,score))
  cutoffs.append({'year':year,'category':category,'schoolCode':r[0],'school':r[1],'majorCode':str(r[2]),'major':r[3],'score':int(score) if score else None,'tie':[int(v) if isinstance(v,(int,float)) else None for v in r[5:12]],'note':str(r[12] or ''),'sourceId':stem,'sourceRow':rownum})

# 仅高校所在城市，不冒充实际就读校区。
schools={};province='';sheet=xlrd.open_workbook(str(ROOT/'data/raw/schools-2026.xls')).sheet_by_index(0)
for i in range(3,sheet.nrows):
 r=sheet.row_values(i)
 if isinstance(r[0],str) and '（' in r[0] and not r[1]:province=r[0].split('（')[0];continue
 if not r[1] or not re.fullmatch(r'\d{10}(?:\.0)?',str(r[2])):continue
 nature='合作办学' if '合作' in r[6] else '民办' if '民办' in r[6] else '公办'
 code=str(int(float(r[2])));schools[r[1]]={'code':code,'province':province,'city':r[4],'level':r[5],'nature':nature,'sponsor':r[3]}
assert len(schools)==2952
add_source('schools-2026','2026年全国普通高等学校名单','https://www.moe.gov.cn/jyb_xxgk/s5743/s5744/202606/W020260618416094865984.xls','schools-2026.xls','2026-06-18','schools','教育部')

# 只接受能识别具体分数且人数/累计相邻相符的分数档，不推测缺失单元格。
scoreSources={2023:'https://file.hebeea.edu.cn/files/article/2023/06/20230624203549_178.pdf',2024:'https://www.centv.cn/a/10001/202406/030fec08d92920c66a1502755d574504.pdf',2025:'https://file.hebeea.edu.cn/files/article/2025/06/20250624193800_658.pdf',2026:'https://file.hebeea.edu.cn/upload/resources/file/2026/06/24/27144.pdf'}
scoreTables={};quality=[]
for year,url in scoreSources.items():
 add_source(f'scores-{year}',f'{year}河北物理/历史考生成绩统计表',url,f'scores-{year}.pdf',f'{year}-06-24','scores')
 raw={}
 for page in range(18):
  cols=[[] for _ in range(5)]
  p=Path(f'/private/tmp/hebei-tess/{year}-{page:02d}.tsv')
  if not p.exists():continue
  for v in csv.DictReader(p.open(),delimiter='\t'):
   t=v['text'].strip().replace(',','');x=(int(v['left'])+int(v['width'])/2)/1587;y=(int(v['top'])+int(v['height'])/2)/2245
   # 图像从官方PDF按原始1587x2245像素提取；仅接受表内数字。
   col=0 if .08<x<.275 else 1 if .28<x<.438 else 2 if .44<x<.60 else 3 if .60<x<.76 else 4 if .76<x<.925 else -1
   if col<0 or y>.95:continue
   match=re.match(r'^(\d+)(.*)$',t)
   if not match or (col>0 and match.group(2)):continue
   num=int(match.group(1))
   if col==0 and not 100<=num<=750:continue
   cols[col].append({'n':num,'y':y,'top':bool(col==0 and match.group(2))})
  for cell in cols[0]:
   score=cell['n']
   if not 100<=score<=750 or cell['y']>.95:continue
   values=[]
   for candidates in cols[1:]:
    nearest=min(candidates,key=lambda a:abs(a['y']-cell['y']),default=None)
    values.append(nearest['n'] if nearest and abs(nearest['y']-cell['y'])<.006 else None)
   if score in raw:raise ValueError(('duplicate_score',year,score))
   raw[score]={'values':values,'page':page+1,'top':cell['top']}
 accepted=collections.Counter();rejected=collections.Counter()
 for category,offset in [('物理',0),('历史',2)]:
  result={}
  for score,r in raw.items():
   count,cum=r['values'][offset:offset+2]
   if count is None or cum is None or count<1 or cum<count:continue
   prev=raw.get(score+1);prevCum=prev['values'][offset+1] if prev else None
   valid=(prevCum is not None and cum-prevCum==count) or (r['top'] and count==cum)
   if valid:result[str(score)]={'low':cum-count+1,'high':cum,'count':count,'page':r['page'],'topBucket':r['top']};accepted[category]+=1
   else:rejected[category]+=1
  scoreTables[f'{year}-{category}']=result
 quality.append({'year':year,'recognizedScores':len(raw),'verified':dict(accepted),'unverified':dict(rejected),'method':'表格OCR；相邻分数累计差与本段人数一致才接受。缺失不补齐。'})

for r in cutoffs:
 t=scoreTables.get(f"{r['year']}-{r['category']}",{}).get(str(r['score']))
 r['rankLow']=t['low'] if t else None;r['rankHigh']=t['high'] if t else None
 r['rankMethod']='同分段区间（非末位个人精确位次）' if t else '扫描表校验未通过或该分数档缺失，位次待核'

# 真实高校分省计划仅作为待审项目样本；缺年度代号/校区/章程核验时禁止冒充完整正式计划。
planURL='https://zsw.jhc.edu.cn/_upload/article/files/a6/44/051647ce4378ac58e4f305b7b445/1ec23b7c-2396-4330-9c35-d957ecfbda38.xlsx'
add_source('plan-jhc-2026','金华职业技术大学2026河北招生计划（高校公布）',planURL,'jhc-plan-2026.xlsx','2026-06-18','plan','金华职业技术大学')
plans=[];w=openpyxl.load_workbook(ROOT/'data/raw/jhc-plan-2026.xlsx',read_only=True,data_only=True)
for i,r in enumerate(w.active.iter_rows(min_row=2,values_only=True),2):
 if r[0]!='河北' or r[1]!='本科' or r[6] not in ['理工/物理类','文史/历史类']:continue
 plans.append({'id':f'jhc2026-{i}','year':2026,'province':'河北','category':'物理' if r[6]=='理工/物理类' else '历史','batch':'本科批','planType':'普通','school':'金华职业技术大学','major':r[2],'schoolCode':'','majorCode':'','seats':int(r[8]),'years':4,'tuition':int(r[5]),'requirement':r[7],'campus':'','source':planURL,'charter':'','published':'2026-06-18','reviewed':False,'reviewNote':'缺当年填报代号、校区及章程核验，不能进入已核验方案。'})
data={'version':'2026-10-29.2','sources':sources,'schools':schools,'cutoffs':cutoffs,'scoreTables':scoreTables,'quality':quality,'planSamples':plans,'coverage':{'cutoffYears':[2023,2024,2025,2026],'planYears':[2026],'planScope':'仅金华职业技术大学高校公布样本；非河北全量正式计划','formal2027':False}}
data['cutoffFiles']=[]
for year in data['coverage']['cutoffYears']:
 filename=f'cutoffs-{year}.json';data['cutoffFiles'].append(filename)
 (ROOT/'dist'/filename).write_text(json.dumps([r for r in cutoffs if r['year']==year],ensure_ascii=False,separators=(',',':')))
del data['cutoffs']
(ROOT/'dist/planning-data.json').write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
print(json.dumps({'cutoffs':len(cutoffs),'rankConverted':sum(r['rankLow'] is not None for r in cutoffs),'schools':len(schools),'quality':quality,'planSamples':len(plans)},ensure_ascii=False))
