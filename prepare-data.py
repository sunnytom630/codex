"""读取公开官方工作簿，保留原始科目要求；不将拟招记录解释为当年计划。"""
import json,collections,hashlib
from pathlib import Path
import openpyxl
source=Path('/private/tmp/hebei-2027-mirror.xlsx')
dest=Path(__file__).parent/'dist'
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
s=w.active
assert s.cell(2,1).value=='2027年拟在河北招生的普通高校招生专业选考科目要求'
schools={};levels=[];requirements=[];records=[];seen=set()
for raw in s.iter_rows(min_row=4,values_only=True):
 if not raw[0] or not raw[1] or not raw[3]:continue
 code,school,pcode,major,level,req=[str(v).strip() if v is not None else '' for v in raw]
 assert level and req,(code,major)
 schools[code]=school
 if level not in levels:levels.append(level)
 if req not in requirements:requirements.append(req)
 record=(code,pcode,major,levels.index(level),requirements.index(req))
 if record in seen:continue
 seen.add(record);records.append(record)
data={'meta':{'title':s.cell(2,1).value,'published':'2025-02-13','checked':'2026-10-29','source':'https://gaokao.chsi.com.cn/gkxx/zc/ss/202502/20250213/2293354258.html','download':'https://t4.chei.com.cn/news/file.do?attach=true&hist=false&id=2293354259&method=downFile','sha256':hashlib.sha256(source.read_bytes()).hexdigest()},'schools':schools,'levels':levels,'requirements':requirements,'records':records}
(dest/'requirements.json').write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
print(json.dumps({'records':len(records),'schools':len(schools),'levels':levels,'requirements':requirements},ensure_ascii=False))
